#!/usr/bin/env python3
"""ネットワーク接続なしで collect.py / build_excel.py の動作を検証する。

食べログのマークアップを模したHTMLを生成し、パース結果と集計値が
期待どおりかを確認する。実データは一切使用しない。

  python3 selftest.py
"""
from __future__ import annotations

import csv
import random
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent

GENRES = ["居酒屋", "ラーメン", "イタリアン", "カフェ", "焼肉", "寿司", "バー", "そば"]
SUBS = ["日本酒バー", "つけ麺", "ピザ", "喫茶店", "ホルモン", "海鮮"]

CARD = """
<div class="list-rst js-rst-cassette-wrap">
  <div class="list-rst__wrap js-open-new-window">
    <h3 class="list-rst__rst-name">
      <a class="list-rst__rst-name-target cpy-rst-name" href="https://tabelog.com/tokyo/A1303/A130301/131{i:05d}/">テスト店舗{i}</a>
    </h3>
    <span class="list-rst__area-genre cpy-area-genre">渋谷 / {genre}</span>
    <span class="c-rating list-rst__rating">
      <span class="c-rating__val c-rating__val--strong list-rst__rating-val">{rating}</span>
    </span>
    <em class="list-rst__rvw-count-num cpy-review-count">{reviews}</em>
    <em class="list-rst__save-count-num">{saves}</em>
    <p class="list-rst__budget">
      <span class="list-rst__budget-val cpy-dinner-budget-val">￥{dinner:,}～</span>
      <span class="list-rst__budget-val cpy-lunch-budget-val">￥{lunch:,}～</span>
    </p>
  </div>
</div>"""

# 店名リンクを持たない広告枠。パース対象から除外されることを確認する。
AD_CARD = '<div class="list-rst"><div class="list-rst__wrap"><span class="ad">広告</span></div></div>'


def make_fixture(path: Path, count: int, unrated_every: int) -> None:
    random.seed(42)
    cards = []
    for i in range(1, count + 1):
        genres = [GENRES[i % len(GENRES)]]
        if i % 2 == 0:
            genres.append(SUBS[i % len(SUBS)])
        rating = "-" if i % unrated_every == 0 else f"{3.0 + (i % 10) / 10:.2f}"
        cards.append(CARD.format(
            i=i, genre="、".join(genres), rating=rating,
            reviews=i * 3, saves=i * 11,
            dinner=3000 + (i % 4) * 2000, lunch=1000 + (i % 3) * 1000,
        ))
    cards.append(AD_CARD)
    path.write_text(
        "<html><body><ul class='js-rstlist-info'>" + "".join(cards) + "</ul></body></html>",
        encoding="utf-8",
    )


def run(cmd: list[str]) -> None:
    result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        print(result.stdout)
        print(result.stderr, file=sys.stderr)
        raise SystemExit(f"コマンド失敗: {' '.join(cmd)}")


def main() -> int:
    count, unrated_every = 40, 7
    expected_unrated = count // unrated_every

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        html = tmp_path / "page1.html"
        csv_path = tmp_path / "out.csv"
        xlsx_path = tmp_path / "out.xlsx"

        make_fixture(html, count, unrated_every)
        run([sys.executable, str(HERE / "collect.py"), "--html", str(html), "-o", str(csv_path)])

        rows = list(csv.DictReader(csv_path.open(encoding="utf-8-sig")))
        checks: list[tuple[str, bool, str]] = []

        checks.append(("広告枠を除いた店舗数", len(rows) == count, f"{len(rows)} (期待 {count})"))

        unrated = sum(1 for r in rows if not r["点数"])
        checks.append(("未評価が空欄", unrated == expected_unrated, f"{unrated} (期待 {expected_unrated})"))

        checks.append(("エリアを分離", all(r["エリア"] == "渋谷" for r in rows), "全行が「渋谷」"))

        multi = [r for r in rows if "、" in r["全ジャンル"]]
        checks.append(("複数ジャンルを保持", len(multi) == count // 2, f"{len(multi)} (期待 {count // 2})"))

        checks.append((
            "主ジャンルは先頭のみ",
            all(r["主ジャンル"] == r["全ジャンル"].split("、")[0] for r in rows),
            "全行一致",
        ))

        checks.append(("URLを取得", all(r["URL"].startswith("https://tabelog.com/") for r in rows), "全行一致"))

        # 主ジャンルの件数合計は総店舗数と一致する（1店1カウント）
        main_total = sum(Counter(r["主ジャンル"] for r in rows).values())
        checks.append(("主ジャンル合計＝総店舗数", main_total == count, f"{main_total} (期待 {count})"))

        run([sys.executable, str(HERE / "build_excel.py"),
             "-i", str(csv_path), "-o", str(xlsx_path), "--source", "セルフテスト（架空データ）"])

        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path)
        expected_sheets = ["サマリー", "統計_点数", "統計_ジャンル", "データ", "ジャンル展開", "凡例・出典"]
        checks.append(("シート構成", wb.sheetnames == expected_sheets, str(wb.sheetnames)))

        ws = wb["データ"]
        checks.append(("データ行数", ws.max_row == count + 1, f"{ws.max_row} (期待 {count + 1})"))

        blank_scores = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=5).value is None)
        checks.append(("未評価は空セル", blank_scores == expected_unrated, f"{blank_scores}"))

        ok = True
        for label, passed, detail in checks:
            print(f"[{'OK ' if passed else 'NG '}] {label}: {detail}")
            ok &= passed

        print("\nセルフテスト:", "成功" if ok else "失敗")
        return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
