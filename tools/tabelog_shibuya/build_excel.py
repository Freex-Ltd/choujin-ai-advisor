#!/usr/bin/env python3
"""collect.py が出力したCSVから、統計シート付きのExcelブックを生成する。

出力シート:
  サマリー      … 全体の件数・平均点など
  統計_点数     … 星（点数）0.1刻みの件数分布 + 棒グラフ
  統計_ジャンル … 飲食店の種類ごとの件数・平均点 + 棒グラフ
  データ        … 収集した生データ
  ジャンル展開  … 1店舗×複数ジャンルを1行1ジャンルに展開（集計用）
  凡例・出典    … 定義と出典の記録

使い方:
  python3 build_excel.py -i shibuya.csv -o 渋谷_食べログ.xlsx --source "食べログ 渋谷 rstLst 2026-08-24取得"
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
from collections import Counter

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
BASE_FONT = Font(name=FONT, size=11)
TITLE_FONT = Font(name=FONT, bold=True, size=14, color="1F3864")
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="666666")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# 点数のビン定義: (ラベル, 下限, 上限) — 上限は「未満」
BINS = [(f"{lo/100:.2f}〜{(lo+9)/100:.2f}", lo / 100, (lo + 10) / 100) for lo in range(300, 400, 10)]


def style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def autosize(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def to_float(value: str):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def to_int(value: str):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def build(rows: list[dict], out_path: str, source: str) -> None:
    wb = Workbook()
    n = len(rows)
    last = n + 1  # データシートの最終行

    # ---------- データ ----------
    ws_data = wb.active
    ws_data.title = "データ"
    headers = ["店名", "主ジャンル", "全ジャンル", "エリア", "点数",
               "口コミ数", "保存件数", "夜予算", "昼予算", "URL"]
    ws_data.append(headers)
    style_header(ws_data, 1, len(headers))

    for row in rows:
        ws_data.append([
            row.get("店名", ""),
            row.get("主ジャンル", ""),
            row.get("全ジャンル", ""),
            row.get("エリア", ""),
            to_float(row.get("点数", "")),      # 未評価は None（空セル）
            to_int(row.get("口コミ数", "")),
            to_int(row.get("保存件数", "")),
            row.get("夜予算", ""),
            row.get("昼予算", ""),
            row.get("URL", ""),
        ])

    for r in range(2, last + 1):
        for c in range(1, len(headers) + 1):
            ws_data.cell(row=r, column=c).font = BASE_FONT
        ws_data.cell(row=r, column=5).number_format = "0.00"
        ws_data.cell(row=r, column=6).number_format = "#,##0"
        ws_data.cell(row=r, column=7).number_format = "#,##0"
    ws_data.freeze_panes = "A2"
    ws_data.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    autosize(ws_data, {1: 34, 2: 16, 3: 30, 4: 14, 5: 8, 6: 10, 7: 10, 8: 18, 9: 18, 10: 46})

    SCORE = f"データ!$E$2:$E${last}"
    NAMES = f"データ!$A$2:$A${last}"

    # ---------- ジャンル展開（全ジャンル集計用の縦持ち） ----------
    ws_long = wb.create_sheet("ジャンル展開")
    ws_long.append(["店名", "ジャンル", "点数"])
    style_header(ws_long, 1, 3)
    long_count = 0
    for row in rows:
        for genre in [g.strip() for g in row.get("全ジャンル", "").split("、") if g.strip()]:
            ws_long.append([row.get("店名", ""), genre, to_float(row.get("点数", ""))])
            long_count += 1
    for r in range(2, long_count + 2):
        for c in (1, 2, 3):
            ws_long.cell(row=r, column=c).font = BASE_FONT
        ws_long.cell(row=r, column=3).number_format = "0.00"
    ws_long.freeze_panes = "A2"
    autosize(ws_long, {1: 34, 2: 20, 3: 8})
    ws_long["E1"] = "このシートは集計用です（1店舗の複数ジャンルを1行1ジャンルに展開）。"
    ws_long["E1"].font = NOTE_FONT
    long_last = long_count + 1
    LONG_GENRE = f"ジャンル展開!$B$2:$B${long_last}"
    LONG_SCORE = f"ジャンル展開!$C$2:$C${long_last}"

    # ---------- 統計_点数 ----------
    ws_score = wb.create_sheet("統計_点数", 0)
    ws_score["A1"] = "星（点数）ごとの件数"
    ws_score["A1"].font = TITLE_FONT
    ws_score.append([])
    ws_score.append(["区分", "件数", "構成比"])
    style_header(ws_score, 3, 3)

    score_rows = []  # (ラベル, 件数フォーミュラ)
    score_rows.append(("未評価（—）", f"=COUNTA({NAMES})-COUNT({SCORE})"))
    score_rows.append(("3.00未満", f'=COUNTIFS({SCORE},">0",{SCORE},"<3")'))
    for label, lo, hi in BINS:
        score_rows.append((label, f'=COUNTIFS({SCORE},">={lo:.2f}",{SCORE},"<{hi:.2f}")'))
    score_rows.append(("4.00以上", f'=COUNTIFS({SCORE},">=4")'))

    first_row = 4
    for offset, (label, formula) in enumerate(score_rows):
        r = first_row + offset
        ws_score.cell(row=r, column=1, value=label).font = BASE_FONT
        ws_score.cell(row=r, column=2, value=formula).font = BASE_FONT
        ws_score.cell(row=r, column=2).number_format = "#,##0"
    last_bin = first_row + len(score_rows) - 1
    total_row = last_bin + 1
    ws_score.cell(row=total_row, column=1, value="合計").font = Font(name=FONT, bold=True, size=11)
    ws_score.cell(row=total_row, column=2, value=f"=SUM(B{first_row}:B{last_bin})").font = Font(name=FONT, bold=True, size=11)
    ws_score.cell(row=total_row, column=2).number_format = "#,##0"

    for r in range(first_row, total_row + 1):
        cell = ws_score.cell(row=r, column=3, value=f'=IF($B${total_row}=0,"—",B{r}/$B${total_row})')
        cell.font = BASE_FONT
        cell.number_format = "0.0%"
        for c in range(1, 4):
            ws_score.cell(row=r, column=c).border = BORDER
    autosize(ws_score, {1: 18, 2: 10, 3: 10})

    note_row = total_row + 2
    ws_score.cell(row=note_row, column=1,
                  value="区分の定義: 下限以上・上限未満の0.1点刻み。「未評価」は食べログ上で点数が「—」の店舗。")
    ws_score.cell(row=note_row, column=1).font = NOTE_FONT
    ws_score.cell(row=note_row + 1, column=1, value=f"出典: {source}")
    ws_score.cell(row=note_row + 1, column=1).font = NOTE_FONT

    chart = BarChart()
    chart.type = "col"
    chart.title = "点数分布（件数）"
    chart.y_axis.title = "件数"
    chart.x_axis.title = "点数区分"
    chart.height, chart.width = 9, 20
    chart.legend = None
    chart.add_data(Reference(ws_score, min_col=2, min_row=3, max_row=last_bin), titles_from_data=True)
    chart.set_categories(Reference(ws_score, min_col=1, min_row=first_row, max_row=last_bin))
    ws_score.add_chart(chart, "E3")

    # ---------- 統計_ジャンル ----------
    ws_genre = wb.create_sheet("統計_ジャンル", 1)
    ws_genre["A1"] = "飲食店の種類（ジャンル）ごとの件数"
    ws_genre["A1"].font = TITLE_FONT
    ws_genre.append([])
    ws_genre.append(["ジャンル", "件数（主ジャンル）", "延べ件数（全ジャンル）", "構成比", "平均点"])
    style_header(ws_genre, 3, 5)

    # 副ジャンルにしか現れないジャンルも表に載せるため、全ジャンルの和集合を対象にする。
    # 並び順はPython側で決定（主ジャンル件数→延べ件数の降順）、値そのものは数式で算出。
    main_counts = Counter(r.get("主ジャンル", "") for r in rows if r.get("主ジャンル", ""))
    all_counts = Counter(
        g.strip()
        for r in rows
        for g in r.get("全ジャンル", "").split("、")
        if g.strip()
    )
    order = sorted(all_counts, key=lambda g: (-main_counts[g], -all_counts[g], g))

    g_first = 4
    for offset, genre in enumerate(order):
        r = g_first + offset
        ws_genre.cell(row=r, column=1, value=genre).font = BASE_FONT
        ws_genre.cell(row=r, column=2, value=f'=COUNTIF(データ!$B$2:$B${last},$A{r})').font = BASE_FONT
        ws_genre.cell(row=r, column=3, value=f'=COUNTIF({LONG_GENRE},$A{r})').font = BASE_FONT
        # 平均点はそのジャンルを掲げる全店（主・副を問わず）を対象にする
        ws_genre.cell(row=r, column=5,
                      value=f'=IFERROR(AVERAGEIFS({LONG_SCORE},{LONG_GENRE},$A{r},{LONG_SCORE},">0"),"—")').font = BASE_FONT
        ws_genre.cell(row=r, column=5).number_format = "0.00"
    g_last = g_first + len(order) - 1 if order else g_first
    g_total = g_last + 1
    ws_genre.cell(row=g_total, column=1, value="合計").font = Font(name=FONT, bold=True, size=11)
    for col in (2, 3):
        cell = ws_genre.cell(row=g_total, column=col, value=f"=SUM({get_column_letter(col)}{g_first}:{get_column_letter(col)}{g_last})")
        cell.font = Font(name=FONT, bold=True, size=11)
        cell.number_format = "#,##0"

    for r in range(g_first, g_total + 1):
        cell = ws_genre.cell(row=r, column=4, value=f'=IF($B${g_total}=0,"—",B{r}/$B${g_total})')
        cell.font = BASE_FONT
        cell.number_format = "0.0%"
        for c in range(1, 6):
            ws_genre.cell(row=r, column=c).border = BORDER
        ws_genre.cell(row=r, column=2).number_format = "#,##0"
        ws_genre.cell(row=r, column=3).number_format = "#,##0"
    autosize(ws_genre, {1: 22, 2: 18, 3: 22, 4: 10, 5: 10})

    gnote = g_total + 2
    ws_genre.cell(row=gnote, column=1, value=(
        "「件数（主ジャンル）」は食べログ表記の先頭ジャンルで1店1カウント（合計＝総店舗数）。"
        "「延べ件数」は複数ジャンル表記をすべて数えるため合計は総店舗数を上回る。"))
    ws_genre.cell(row=gnote, column=1).font = NOTE_FONT
    ws_genre.cell(row=gnote + 1, column=1, value=(
        "副ジャンルにのみ現れるジャンルも行として掲載する（その場合「件数（主ジャンル）」は0）。"
        "平均点はそのジャンルを掲げる全店が対象で、未評価店は除外。"))
    ws_genre.cell(row=gnote + 1, column=1).font = NOTE_FONT
    ws_genre.cell(row=gnote + 2, column=1, value=f"出典: {source}")
    ws_genre.cell(row=gnote + 2, column=1).font = NOTE_FONT

    top_last = min(g_first + 14, g_last)  # 上位15ジャンルをグラフ化
    if order:
        gchart = BarChart()
        gchart.type = "bar"
        gchart.title = "ジャンル別 店舗数（上位15）"
        gchart.x_axis.title = "件数"
        gchart.height, gchart.width = 12, 20
        gchart.legend = None
        gchart.add_data(Reference(ws_genre, min_col=2, min_row=3, max_row=top_last), titles_from_data=True)
        gchart.set_categories(Reference(ws_genre, min_col=1, min_row=g_first, max_row=top_last))
        ws_genre.add_chart(gchart, "G3")

    # ---------- サマリー ----------
    ws_sum = wb.create_sheet("サマリー", 0)
    ws_sum["A1"] = "渋谷 飲食店データ サマリー"
    ws_sum["A1"].font = Font(name=FONT, bold=True, size=16, color="1F3864")
    ws_sum["A2"] = f"出典: {source}"
    ws_sum["A2"].font = NOTE_FONT
    ws_sum["A3"] = f"ブック生成日: {dt.date.today().isoformat()}"
    ws_sum["A3"].font = NOTE_FONT

    ws_sum.append([])
    ws_sum.append(["指標", "値"])
    style_header(ws_sum, 5, 2)
    metrics = [
        ("総店舗数", f"=COUNTA({NAMES})", "#,##0"),
        ("点数あり店舗数", f"=COUNT({SCORE})", "#,##0"),
        ("未評価店舗数", f"=COUNTA({NAMES})-COUNT({SCORE})", "#,##0"),
        ("平均点（未評価除く）", f'=IFERROR(AVERAGEIF({SCORE},">0"),"—")', "0.00"),
        ("中央値（未評価除く）", f'=IFERROR(MEDIAN({SCORE}),"—")', "0.00"),
        ("最高点", f'=IFERROR(MAX({SCORE}),"—")', "0.00"),
        ("最低点（未評価除く）", f'=IFERROR(MIN({SCORE}),"—")', "0.00"),
        ("ジャンル数（全ジャンル）", f"=COUNTA(統計_ジャンル!$A${g_first}:$A${g_last})", "#,##0"),
        ("ジャンル数（主ジャンル）", f'=COUNTIF(統計_ジャンル!$B${g_first}:$B${g_last},">0")', "#,##0"),
        ("口コミ総数", f"=SUM(データ!$F$2:$F${last})", "#,##0"),
    ]
    for offset, (label, formula, fmt) in enumerate(metrics):
        r = 6 + offset
        ws_sum.cell(row=r, column=1, value=label).font = BASE_FONT
        cell = ws_sum.cell(row=r, column=2, value=formula)
        cell.font = BASE_FONT
        cell.number_format = fmt
        for c in (1, 2):
            ws_sum.cell(row=r, column=c).border = BORDER
    autosize(ws_sum, {1: 30, 2: 16})

    # ---------- 凡例・出典 ----------
    ws_doc = wb.create_sheet("凡例・出典")
    ws_doc["A1"] = "凡例・出典・前提"
    ws_doc["A1"].font = TITLE_FONT
    notes = [
        ("出典", source),
        ("収集方法", "食べログのエリア一覧ページ（rstLst）のHTMLを collect.py で解析"),
        ("点数", "食べログの表示点数をそのまま転記。未評価（—）は空セル"),
        ("主ジャンル", "食べログのジャンル表記の先頭1つ。1店舗1カウントで合計＝総店舗数"),
        ("全ジャンル", "食べログのジャンル表記すべて。延べ件数の集計に使用"),
        ("ジャンル一覧", "主ジャンル・副ジャンルの和集合。副ジャンルにのみ現れるジャンルは「件数（主ジャンル）」が0になる"),
        ("点数区分", "0.1点刻み（下限以上・上限未満）。3.00未満と4.00以上は別枠"),
        ("平均点", "未評価店を除外して算出。ジャンル別平均はそのジャンルを掲げる全店（主・副を問わず）が対象"),
        ("注意", "食べログの点数は独自アルゴリズムによるもので、取得時点の値。時間経過で変動する"),
        ("注意", "一覧ページに表示されない店舗（掲載終了・非公開など）は含まれない"),
    ]
    ws_doc.append([])
    ws_doc.append(["項目", "内容"])
    style_header(ws_doc, 3, 2)
    for offset, (key, value) in enumerate(notes):
        r = 4 + offset
        ws_doc.cell(row=r, column=1, value=key).font = BASE_FONT
        cell = ws_doc.cell(row=r, column=2, value=value)
        cell.font = BASE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        for c in (1, 2):
            ws_doc.cell(row=r, column=c).border = BORDER
    autosize(ws_doc, {1: 16, 2: 82})

    wb.save(out_path)
    print(f"{out_path} を生成しました（{n}店舗 / ジャンル{len(order)}種）。")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-i", "--input", required=True, help="collect.py が出力したCSV")
    ap.add_argument("-o", "--out", default="渋谷_食べログ.xlsx", help="出力xlsxパス")
    ap.add_argument("--source", default="食べログ（取得日未記入）", help="出典として記録する文字列")
    args = ap.parse_args()

    with open(args.input, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        raise SystemExit("CSVにデータがありません。")
    build(rows, args.out, args.source)


if __name__ == "__main__":
    main()
